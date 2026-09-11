import importlib.util
import io
import uuid
from pathlib import Path

import pytest
from app import models
from app.database import SessionLocal, get_session
from app.main import app
from app.portfolio_seed import DEMO_SOURCE
from app.terminal_analytics import history
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

client = TestClient(app)


def test_portfolio_operating_api_and_export() -> None:
    response = client.get("/api/v1/operations/portfolio")
    assert response.status_code == 200
    data = response.json()
    assert data["portfolio"]["code"] == "KNK_MAIN"
    assert data["reconciliation"]["state"] == "BALANCED"
    assert len(data["metric_metadata"]) > 20
    assert client.get("/api/v1/operations/trades").json()["items"]
    export = client.get("/api/v1/operations/export")
    assert export.status_code == 200, export.text[:200]
    book = load_workbook(io.BytesIO(export.content), data_only=True)
    assert {"Summary", "Positions", "NAV History", "Trade Monitor", "Metadata"} <= set(
        book.sheetnames
    )
    assert (
        client.post("/api/v1/operations/reset-demo", json={"confirmation": "no"}).status_code == 422
    )
    assert (
        client.post("/api/v1/operations/reconciliation").json()["state"] == "BROKER_NOT_CONNECTED"
    )


def test_agent_pairing_scopes_single_use_revocation() -> None:
    pairing = client.post("/api/v1/data-drop/agents/pair").json()
    claim = client.post(
        "/agent/v1/pair", json={"code": pairing["code"], "name": "Test agent " + str(uuid.uuid4())}
    )
    assert claim.status_code == 200
    payload = claim.json()
    headers = {"Authorization": "Bearer " + payload["token"]}
    assert payload["scopes"] == ["files:upload", "files:status", "agent:heartbeat"]
    assert client.post(
        "/agent/v1/pair", json={"code": pairing["code"], "name": "Replay"}
    ).status_code in {403, 409}
    assert (
        client.post("/agent/v1/heartbeat", json={"queued": 1}, headers=headers).status_code == 200
    )
    assert client.get("/agent/v1/files").status_code == 401
    upload = client.post(
        "/agent/v1/files",
        files={"file": (f"agent-{uuid.uuid4()}.csv", "date,close\n2026-09-04,100\n", "text/csv")},
        headers=headers,
    )
    assert upload.status_code == 200
    files = client.get("/agent/v1/files", headers=headers).json()["items"]
    assert upload.json()["id"] in {f["id"] for f in files}
    assert "token" not in client.get("/api/v1/data-drop/agents").text
    assert (
        client.post("/api/v1/data-drop/agents/" + payload["agent_id"] + "/revoke").status_code
        == 200
    )
    assert client.get("/agent/v1/files", headers=headers).status_code == 401


def test_agent_endpoint_and_path_boundaries(tmp_path: Path) -> None:
    spec = importlib.util.spec_from_file_location(
        "local_data_agent", Path(__file__).parents[2] / "local-agent" / "agent.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    import pytest

    with pytest.raises(ValueError):
        module.validate_endpoint("http://example.com", True)
    with pytest.raises(ValueError):
        module.validate_endpoint("https://token@example.com", False)
    with pytest.raises(ValueError):
        module.safe_path(tmp_path, tmp_path / ".." / "outside.csv")
    assert module.validate_endpoint("http://127.0.0.1:8000", True) == "http://127.0.0.1:8000"
    if module.sys.platform == "win32":
        assert type(module.credential_backend()).__name__ == "WinVaultKeyring"


def test_candidate_requires_pinned_version_and_blocks_unverified_promotion() -> None:
    from app import models
    from app.database import SessionLocal
    from sqlalchemy import select

    with SessionLocal() as session:
        version = session.scalar(select(models.DatasetVersion))
        assert version is not None
        version_id = version.id
    created = client.post(
        "/api/v1/desks/candidates",
        json={
            "name": "Research " + str(uuid.uuid4()),
            "hypothesis": "Trailing momentum hypothesis; no proven edge asserted.",
            "dataset_version_id": version_id,
        },
    )
    assert created.status_code == 200, created.text
    data = created.json()
    assert data["state"] == "RESEARCH"
    assert data["review"]["gates"]["paper_validation"] == "NOT CONNECTED"
    denied = client.post(
        "/api/v1/desks/candidates/" + data["id"] + "/review",
        json={
            "state": "PAPER TESTING",
            "note": "Review without verified paper history",
        },
    )
    assert denied.status_code == 422
    reviewed = client.post(
        "/api/v1/desks/candidates/" + data["id"] + "/review",
        json={
            "state": "REJECTED",
            "note": "Insufficient evidence after review",
        },
    )
    assert reviewed.status_code == 200 and reviewed.json()["state"] == "REJECTED"
    missing = client.post(
        "/api/v1/desks/candidates",
        json={
            "name": "Invalid candidate",
            "hypothesis": "Missing immutable dataset history",
            "dataset_version_id": "missing",
        },
    )
    assert missing.status_code == 422


def test_factor_horizon_never_silently_shortens(monkeypatch: pytest.MonkeyPatch) -> None:
    # Pin the deliberately short managed series; legacy history now remains
    # available unless the selected source explicitly excludes it.
    with SessionLocal() as session:
        for item in session.scalars(select(models.Instrument)).all():
            rule = session.scalar(
                select(models.SourcePrecedenceRule).where(
                    models.SourcePrecedenceRule.instrument_id == item.id
                )
            )
            if rule is None:
                rule = models.SourcePrecedenceRule(
                    instrument_id=item.id,
                    priority=["BROKER", "PROVIDER", "FILE", "DEMO"],
                    reason="Short factor-history test fixture",
                )
                session.add(rule)
            rule.preferred_source = DEMO_SOURCE
        session.flush()
        monkeypatch.setitem(app.dependency_overrides, get_session, lambda: session)
        selected = history(session, "AAPL", 800)["items"]
        assert 5 < len(selected) < 252
        assert all(row.get("source") == DEMO_SOURCE for row in selected)
        response = client.get("/api/v1/factors?lookback=252")
        assert response.status_code == 200
        assert response.json()["quality"] == "INSUFFICIENT DATA"
        assert response.json()["lookback"] == 252
        assert response.json()["items"] == []
