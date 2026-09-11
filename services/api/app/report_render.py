"""Portable report adapters. Render exclusively from a validated pinned snapshot."""

from decimal import Decimal, InvalidOperation
from io import BytesIO
from xml.sax.saxutils import escape

from .report_contracts import Cell, ReportSection, ReportSnapshot


def display(value: Cell, limit: int = 120) -> str:
    text = "Unavailable" if value is None else str(value)
    return text if len(text) <= limit else text[: limit - 3] + "..."


def metadata(snapshot: ReportSnapshot) -> ReportSection:
    result = ReportSection(
        title="Sources",
        columns=["Field", "Value"],
        rows=[
            ["Report", snapshot.title],
            ["Generated UTC", snapshot.requested_at],
            ["Data timestamp", snapshot.data_as_of],
            ["Source", snapshot.source],
            ["State", snapshot.quality],
            ["Currency", snapshot.currency],
            ["Calculation version", snapshot.calculation_version],
            ["Disclosure", snapshot.disclosure],
            [
                "Snapshot",
                "Complete immutable input is available from the authenticated job source download",
            ],
        ],
    )
    for key, value in snapshot.references.items():
        result.rows.append([key, str(value) if value is not None else None])
    return result


def xlsx(snapshot: ReportSnapshot) -> bytes:
    import xlsxwriter  # type: ignore[import-untyped]  # 3.2.0 has no published stubs.

    output = BytesIO()
    workbook = xlsxwriter.Workbook(
        output, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}
    )
    workbook.set_properties(
        {"title": snapshot.title, "company": "KnK Capital", "comments": snapshot.disclosure}
    )
    header = workbook.add_format(
        {"bold": True, "bg_color": "#202124", "font_color": "#F5B642", "text_wrap": True}
    )
    body = workbook.add_format({"font_color": "#166534", "text_wrap": True, "valign": "top"})
    number = workbook.add_format(
        {"num_format": '#,##0.00;[Red](#,##0.00);"-"', "font_color": "#166534"}
    )
    formula = workbook.add_format(
        {"num_format": '#,##0.00;[Red](#,##0.00);"-"', "font_color": "#111111"}
    )
    charts = None
    chart_count = 0
    for index, item in enumerate([metadata(snapshot), *snapshot.sections]):
        name = f"{index:02d} {item.title}"[:31]
        sheet = workbook.add_worksheet(name)
        sheet.hide_gridlines(2)
        sheet.freeze_panes(1, 1)
        sheet.set_landscape()
        sheet.fit_to_pages(1, 0)
        sheet.repeat_rows(0)
        sheet.set_header("&LKnK Capital&RPrivate internal research")
        sheet.set_footer("&L" + snapshot.quality + "&RPage &P of &N")
        sheet.set_column(0, max(0, len(item.columns) - 1), 22)
        sheet.set_row(0, 32)
        sheet.write_row(0, 0, item.columns, header)
        for row_index, row in enumerate(item.rows, 1):
            for column, value in enumerate(row):
                sheet.write(
                    row_index,
                    column,
                    value,
                    number
                    if isinstance(value, (int, float)) and not isinstance(value, bool)
                    else body,
                )
        if item.rows and item.columns:
            sheet.autofilter(0, 0, len(item.rows), len(item.columns) - 1)
        # Formula inputs retain original values. Missing fields never become zero.
        fields = ["quantity", "market_price", "contract_multiplier", "fx_rate"]
        if item.title == "Positions" and all(field in item.columns for field in fields):
            target = len(item.columns)
            sheet.write(0, target, f"Calculated value ({snapshot.currency})", header)
            for row_index, row in enumerate(item.rows, 1):
                values = [row[item.columns.index(field)] for field in fields]
                try:
                    decimals = [Decimal(str(value)) for value in values]
                    if not all(value.is_finite() for value in decimals):
                        continue
                    product = decimals[0] * decimals[1] * decimals[2] * decimals[3]
                except InvalidOperation:
                    continue
                refs = [
                    xlsxwriter.utility.xl_rowcol_to_cell(row_index, item.columns.index(field))
                    for field in fields
                ]
                sheet.write_formula(
                    row_index, target, "=" + "*".join(refs), formula, float(product)
                )
        value_key = next((key for key in ("nav", "equity") if key in item.columns), None)
        if value_key and "date" in item.columns and item.rows:
            value_index = item.columns.index(value_key)
            date_index = item.columns.index("date")
            try:
                chart_values = [Decimal(str(row[value_index])) for row in item.rows]
                if not all(value.is_finite() for value in chart_values):
                    continue
            except InvalidOperation:
                continue
            if charts is None:
                charts = workbook.add_worksheet("Charts")
                charts.hide_gridlines(2)
                charts.write(0, 0, snapshot.title, header)
                charts.set_column(0, 10, 12)
            chart = workbook.add_chart({"type": "line"})
            chart.add_series(
                {
                    "name": f"{value_key.upper()} ({snapshot.currency})",
                    "categories": [name, 1, date_index, len(item.rows), date_index],
                    "values": [name, 1, value_index, len(item.rows), value_index],
                    "values_data": [float(value) for value in chart_values],
                    "line": {"color": "#137E77", "width": 1.5},
                }
            )
            chart.set_title({"name": item.title})
            chart.set_legend({"none": True})
            chart.set_size({"width": 800, "height": 340})
            charts.insert_chart(2 + chart_count * 19, 0, chart)
            chart_count += 1
    workbook.close()
    return output.getvalue()


def pptx(snapshot: ReportSnapshot) -> bytes:
    from pptx import Presentation
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = snapshot.title
    deck.core_properties.author = "KnK Capital"
    for section in [metadata(snapshot), *snapshot.sections]:
        # Chunk wide/long tables instead of shrinking text until it is unreadable.
        for column_start in range(0, len(section.columns), 5):
            columns = section.columns[column_start : column_start + 5]
            rows = section.rows or [[None] * len(section.columns)]
            for start in range(0, len(rows), 8):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                title = slide.shapes.add_textbox(
                    Inches(0.45), Inches(0.3), Inches(12.3), Inches(0.6)
                )
                title.text_frame.text = f"KnK Capital | {section.title}"
                title.text_frame.paragraphs[0].font.size = Pt(26)
                title.text_frame.paragraphs[0].font.color.theme_color = MSO_THEME_COLOR.DARK_1
                part = rows[start : start + 8]
                table = slide.shapes.add_table(
                    len(part) + 1,
                    len(columns),
                    Inches(0.45),
                    Inches(1.05),
                    Inches(12.3),
                    Inches(5.45),
                ).table
                table_rows: list[list[Cell]] = [[cell for cell in columns]]
                table_rows.extend(row[column_start : column_start + 5] for row in part)
                for r, values in enumerate(table_rows):
                    for c, value in enumerate(values):
                        cell = table.cell(r, c)
                        cell.text = display(value, 95)
                        cell.fill.solid()
                        cell.fill.fore_color.theme_color = (
                            MSO_THEME_COLOR.DARK_1 if r == 0 else MSO_THEME_COLOR.LIGHT_1
                        )
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.size = Pt(12)
                            paragraph.font.color.theme_color = (
                                MSO_THEME_COLOR.LIGHT_1 if r == 0 else MSO_THEME_COLOR.DARK_1
                            )
                footer = slide.shapes.add_textbox(
                    Inches(0.45), Inches(6.7), Inches(12.3), Inches(0.5)
                )
                footer.text_frame.text = f"{snapshot.quality} | {snapshot.data_as_of or 'Timestamp unavailable'} | {snapshot.currency}\nPrivate internal research | Full unabridged values in job source snapshot | {len(deck.slides)}"
                for paragraph in footer.text_frame.paragraphs:
                    paragraph.font.size = Pt(10)
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def pdf(snapshot: ReportSnapshot) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Flowable,
        LongTable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
        title=snapshot.title,
        author="KnK Capital",
    )
    styles = getSampleStyleSheet()
    styles["BodyText"].fontSize = 8
    styles["BodyText"].leading = 10
    styles["BodyText"].wordWrap = "CJK"
    story: list[Flowable] = [Paragraph(escape(snapshot.title), styles["Title"])]
    for item in [metadata(snapshot), *snapshot.sections]:
        story.append(Paragraph(escape(item.title), styles["Heading2"]))
        for start in range(0, len(item.columns), 5):
            columns = item.columns[start : start + 5]
            rows: list[list[Cell]] = [[cell for cell in columns]]
            rows.extend(row[start : start + 5] for row in item.rows)
            cells = [
                [Paragraph(escape(display(cell, 500)), styles["BodyText"]) for cell in row]
                for row in rows
            ]
            table = LongTable(
                cells,
                colWidths=[document.width / len(columns)] * len(columns),
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F5B642")),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F1F3F4")],
                        ),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.extend([table, Spacer(1, 12)])
    document.build(story)
    return output.getvalue()


def render(snapshot: ReportSnapshot, format: str) -> bytes:
    if format == "xlsx":
        return xlsx(snapshot)
    if format == "pptx":
        return pptx(snapshot)
    if format == "pdf":
        return pdf(snapshot)
    raise ValueError("Unsupported report format")
