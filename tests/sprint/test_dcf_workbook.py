from copy import deepcopy
from decimal import Decimal
from io import BytesIO

import pytest
from app.dcf_workbook import DcfInput, Formula, model_sheets
from app.equity_valuation import DcfRequest, DcfScenario, WaccInputs, dcf_scenarios
from app.report_contracts import ReportSnapshot
from app.report_render import render
from openpyxl import load_workbook
from pydantic import JsonValue


def dcf_snapshot(
    request: DcfRequest | None = None, price: Decimal | None = Decimal(12)
) -> ReportSnapshot:
    request = request or DcfRequest(
        symbol="AAA",
        initial_working_capital=Decimal(10),
        scenarios=[
            DcfScenario(
                name="BASE",
                revenue_growth=[Decimal(".1")],
                ebit_margins=[Decimal(".2")],
                tax_rate=Decimal(".25"),
                terminal_growth=Decimal(0),
            )
        ],
    )
    base: dict[str, JsonValue] = {
        "year": "2025",
        "revenue": 100,
        "debt": 20,
        "cash": 10,
        "shares": 10,
    }
    payload = dcf_scenarios(base, request, price)
    payload.update(
        parameters=request.model_dump(mode="json"),
        quote={"price": str(price) if price is not None else None},
        unit="millions",
    )
    return ReportSnapshot(
        title="KnK Capital | DCF Review",
        kind="dcf",
        requested_at="2026-01-10T00:00:00Z",
        data_as_of="2026-01-09T00:00:00Z",
        source="Fictional statements",
        quality="DEMO",
        currency="USD",
        calculation_version="knk-fcff-1.0",
        references={"analysis_run_id": "pinned-fixture"},
        sections=[],
        payload=payload,
    )


def test_native_forecast_formulas_and_cached_financial_reconciliation() -> None:
    snapshot = dcf_snapshot()
    original = deepcopy(snapshot)
    content = render(snapshot, "xlsx")
    formulas = load_workbook(BytesIO(content), data_only=False)
    cached = load_workbook(BytesIO(content), data_only=True)
    assert formulas["DCF BASE"]["C10"].value == "=B10*(1+C9)"
    assert formulas["DCF BASE"]["C18"].value == "=C17-B17"
    assert formulas["DCF BASE"]["C13"].value == "=MAX(0,C12)*$B$5"
    assert formulas["DCF BASE"]["C19"].value == "=C14+C15-C16-C18"
    assert cached["DCF BASE"]["C19"].value == pytest.approx(13.3)
    assert cached["DCF BASE"]["B26"].value == pytest.approx(133)
    assert cached["DCF Summary"]["B2"].value == pytest.approx(12.3)
    assert cached["DCF Summary"]["H2"].value == 0
    assert len(formulas["DCF BASE"].data_validations.dataValidation) == 8
    assert snapshot == original


@pytest.mark.parametrize("name", ["BASE", "BULL", "BEAR"])
def test_all_sensitivity_cells_link_to_the_forecast(name: str) -> None:
    request = DcfRequest.model_validate({"symbol": "AAA", "scenarios": [{"name": name}]})
    snapshot = dcf_snapshot(request)
    saved = DcfInput.model_validate(snapshot.payload)
    content = render(snapshot, "xlsx")
    formulas = load_workbook(BytesIO(content), data_only=False)
    cached = load_workbook(BytesIO(content), data_only=True)
    sheet = f"DCF {name} Sens"
    for start, cases in (
        (3, saved.scenarios[0].growth_sensitivity),
        (31, saved.scenarios[0].exit_sensitivity),
    ):
        for index, case in enumerate(cases, start):
            assert "SUMPRODUCT" in str(formulas[sheet][f"C{index}"].value)
            assert f"'DCF {name}'!C19:G19" in str(formulas[sheet][f"C{index}"].value)
            assert (
                cached[sheet][f"C{index}"].value == pytest.approx(float(case.fair_value))
                if case.fair_value is not None
                else cached[sheet][f"C{index}"].value == "#N/A"
            )


def test_calculated_wacc_is_a_linked_model_not_a_hardcoded_rate() -> None:
    request = DcfRequest(
        symbol="AAA",
        wacc_inputs=WaccInputs(),
        apply_calculated_wacc=True,
        scenarios=[DcfScenario(name="BASE")],
    )
    content = render(dcf_snapshot(request), "xlsx")
    formulas = load_workbook(BytesIO(content), data_only=False)
    cached = load_workbook(BytesIO(content), data_only=True)
    assert formulas["DCF BASE"]["B6"].value == "='DCF WACC'!B15"
    assert formulas["DCF WACC"]["B15"].value == "=B13*B11+B14*B12"
    assert cached["DCF WACC"]["B15"].value == pytest.approx(0.0719)


def test_missing_quote_and_invalid_sensitivity_are_not_cached_as_zero() -> None:
    request = DcfRequest(
        symbol="AAA",
        scenarios=[DcfScenario(name="BASE", wacc=Decimal(".03"), terminal_growth=Decimal(".025"))],
    )
    book = load_workbook(BytesIO(render(dcf_snapshot(request, None), "xlsx")), data_only=True)
    assert book["DCF Inputs"]["B11"].value is None
    assert book["DCF BASE"]["B32"].value == "#N/A"
    assert book["DCF BASE"]["B33"].value == "#N/A"
    assert book["DCF BASE Sens"]["C3"].value == "#N/A"


@pytest.mark.parametrize("field", ["fair_value", "terminal_value", "forecast", "assumptions"])
def test_inconsistent_saved_result_is_rejected_before_export(field: str) -> None:
    snapshot = dcf_snapshot()
    scenarios = snapshot.payload["scenarios"]
    assert isinstance(scenarios, list) and isinstance(scenarios[0], dict)
    scenarios[0][field] = "0"
    with pytest.raises(ValueError):
        render(snapshot, "xlsx")


def test_no_external_formula_references_or_user_string_interpolation() -> None:
    snapshot = dcf_snapshot()
    snapshot.title = '=HYPERLINK("https://invalid.test", "bad")'
    content = render(snapshot, "xlsx")
    book = load_workbook(BytesIO(content), data_only=False)
    assert book["DCF Inputs"]["A1"].data_type == "s"
    for sheet in model_sheets(snapshot):
        for cell in sheet.cells:
            if isinstance(cell.value, Formula):
                assert "http" not in cell.value.expression
                assert "[" not in cell.value.expression


def test_mismatched_calculation_metadata_cannot_label_a_different_model() -> None:
    snapshot = dcf_snapshot()
    snapshot.calculation_version = "unknown-model"
    with pytest.raises(ValueError, match="metadata must match"):
        render(snapshot, "xlsx")
