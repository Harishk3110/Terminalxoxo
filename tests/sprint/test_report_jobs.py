import hashlib
from datetime import UTC, datetime, timedelta
from io import BytesIO
from zipfile import ZipFile

import pytest
from app import models, report_api, report_jobs
from app.auth_sessions import token_digest
from app.database import get_session
from app.object_storage import ObjectStorage
from app.report_contracts import FORMATS, ReportRequest, ReportSection, ReportSnapshot, digest
from app.report_models import ReportJob
from app.report_render import render
from app.report_sources import capture
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker


@pytest.fixture
def snapshot():
    return ReportSnapshot(
        title="KnK Capital | Portfolio Review",
        kind="portfolio",
        requested_at=datetime.now(UTC).isoformat(),
        data_as_of="2026-01-09T20:00:00+00:00",
        source="Fictional fixture",
        quality="FILE IMPORT",
        currency="SGD",
        calculation_version="fixture-1",
        references={"dataset_version_id": "original-version"},
        sections=[
            ReportSection(
                title="Positions",
                columns=["symbol", "quantity", "market_price", "contract_multiplier", "fx_rate"],
                rows=[["AAA", "2", "3", "100", "1.3"], ["MISSING", 2, None, 100, 1.3]],
            )
        ],
        payload={"nav": "780", "dataset_version_id": "original-version"},
    )


@pytest.fixture
def reports(ledger_session, monkeypatch, tmp_path):
    from app.config import get_settings

    settings = get_settings().model_copy(
        update={
            "object_storage_local_dir": str(tmp_path),
            "object_storage_endpoint": None,
            "object_storage_access_key": None,
            "object_storage_secret_key": None,
        }
    )
    monkeypatch.setattr("app.object_storage.get_settings", lambda: settings)
    factory = sessionmaker(ledger_session.bind, expire_on_commit=False)
    monkeypatch.setattr(report_jobs, "SessionLocal", factory)
    for identifier in ("report-owner", "report-other"):
        ledger_session.add(
            models.User(
                id=identifier,
                email=identifier + "@example.test",
                password_hash="session-only",
                role="ADMIN",
            )
        )
        ledger_session.flush()
        ledger_session.add(
            models.UserSession(
                user_id=identifier,
                session_hash=token_digest(identifier + "-token"),
                expires_at=datetime.now(UTC) + timedelta(hours=1),
            )
        )
    ledger_session.commit()
    app = FastAPI()
    app.include_router(report_api.router)

    def session_override():
        with factory() as session:
            yield session

    app.dependency_overrides[get_session] = session_override
    client = TestClient(app)
    client.cookies.set("knk_session", "report-owner-token")
    return client, factory


@pytest.mark.parametrize(
    "kind,format", [(kind, format) for kind, formats in FORMATS.items() for format in formats]
)
def test_all_sixteen_report_outputs_are_real_files(snapshot, kind, format):
    document = snapshot.model_copy(update={"kind": kind})
    content = render(document, format)
    assert len(content) > 1000
    if format == "pdf":
        assert content.startswith(b"%PDF-")
        assert b"%%EOF" in content
    else:
        with ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()
            assert "[Content_Types].xml" in names
            assert ("xl/workbook.xml" if format == "xlsx" else "ppt/presentation.xml") in names


def test_workbook_formulas_match_source_and_leave_missing_unavailable(snapshot):
    content = render(snapshot, "xlsx")
    formulas = load_workbook(BytesIO(content), data_only=False)
    cached = load_workbook(BytesIO(content), data_only=True)
    assert formulas["01 Positions"]["F2"].value == "=B2*C2*D2*E2"
    assert cached["01 Positions"]["F2"].value == pytest.approx(780)
    assert cached["01 Positions"]["F3"].value is None
    assert formulas["00 Sources"]["B6"].value == "FILE IMPORT"


def test_spreadsheet_content_cannot_inject_formula(snapshot):
    snapshot.sections[0].rows[0][0] = '=HYPERLINK("https://invalid.test","bad")'
    book = load_workbook(BytesIO(render(snapshot, "xlsx")))
    assert book["01 Positions"]["A2"].data_type == "s"


def test_workbook_chart_uses_pinned_unrounded_curve(snapshot):
    snapshot.sections.append(
        ReportSection(
            title="Curve",
            columns=["date", "nav"],
            rows=[["2026-01-01", "100000"], ["2026-01-02", "100123.4567"]],
        )
    )
    with ZipFile(BytesIO(render(snapshot, "xlsx"))) as archive:
        chart = archive.read("xl/charts/chart1.xml").decode()
    assert "100123.4567" in chart
    assert "02 Curve" in chart


def test_local_write_once_does_not_replace_existing(reports):
    storage = ObjectStorage()
    storage.put_new_bytes(key="private-reports/test.txt", data=b"first", content_type="text/plain")
    with pytest.raises(FileExistsError):
        storage.put_new_bytes(
            key="private-reports/test.txt", data=b"changed", content_type="text/plain"
        )
    assert storage.get_bytes("private-reports/test.txt") == b"first"


def test_s3_write_once_uses_atomic_precondition(reports, monkeypatch):
    import boto3
    from botocore.stub import Stubber

    client = boto3.client("s3", aws_access_key_id="fixture", aws_secret_access_key="fixture")
    storage = ObjectStorage()
    monkeypatch.setattr(storage, "_s3", client)
    with Stubber(client) as stub:
        stub.add_response(
            "put_object",
            {},
            {
                "Bucket": storage.settings.object_storage_bucket,
                "Key": "private-reports/test",
                "Body": b"body",
                "ContentType": "text/plain",
                "IfNoneMatch": "*",
            },
        )
        stored = storage.put_new_bytes(
            key="private-reports/test", data=b"body", content_type="text/plain"
        )
        assert stored.content_hash == hashlib.sha256(b"body").hexdigest()
        stub.assert_no_pending_responses()


def test_report_health_requires_real_dependencies_and_recent_poll(reports, monkeypatch):
    import time

    from app import report_engine

    _, factory = reports
    monkeypatch.setattr(report_engine, "SessionLocal", factory)
    monkeypatch.setattr(report_engine, "poll_started", 0)
    monkeypatch.setattr(report_engine, "last_poll", 0)
    client = TestClient(report_engine.app)
    assert client.get("/health/live").status_code == 200
    assert client.get("/health/ready").status_code == 503
    monkeypatch.setattr(report_engine, "last_poll", time.monotonic())
    assert client.get("/health/ready").status_code == 200
    metrics = client.get("/metrics")
    assert metrics.status_code == 200
    assert 'knk_report_jobs{state="QUEUED"} 0.0' in metrics.text
    assert "knk_report_dispatcher_ready 1.0" in metrics.text
    monkeypatch.setattr(ObjectStorage, "get_bytes", lambda *_: b"incorrect")
    assert client.get("/health/ready").status_code == 503


def test_busy_worker_readiness_is_bounded(monkeypatch: pytest.MonkeyPatch) -> None:
    import time

    from app import report_engine

    monkeypatch.setattr(report_engine, "last_poll", 0)
    monkeypatch.setattr(report_engine, "poll_started", time.monotonic() - 30)
    assert report_engine.dispatcher_ready()
    monkeypatch.setattr(report_engine, "poll_started", time.monotonic() - 901)
    assert not report_engine.dispatcher_ready()


@pytest.mark.parametrize("format", ["xlsx", "pptx", "pdf"])
def test_queue_worker_owner_download_and_durable_snapshot(reports, snapshot, monkeypatch, format):
    client, factory = reports
    monkeypatch.setattr(report_api, "capture", lambda *_: snapshot)
    response = client.post("/api/v1/report-jobs", json={"kind": "portfolio", "format": format})
    assert response.status_code == 202
    identifier = response.json()["id"]
    assert response.json()["download_url"] is None
    assert report_jobs.execute(identifier)
    assert not report_jobs.execute(identifier)
    with factory() as session:
        job = session.get(ReportJob, identifier)
        assert job.status == "SUCCEEDED"
        assert job.snapshot["references"]["dataset_version_id"] == "original-version"
        assert job.snapshot_hash == digest(job.snapshot)
        assert job.worker_id
    result = client.get(f"/api/v1/report-jobs/{identifier}").json()
    download = client.get(result["download_url"])
    assert download.status_code == 200
    assert hashlib.sha256(download.content).hexdigest() == result["content_hash"]
    assert "no-store" in download.headers["cache-control"]
    assert "object_key" not in result
    client.cookies.set("knk_session", "report-other-token")
    for suffix in ("", "/download", "/source"):
        assert client.get(f"/api/v1/report-jobs/{identifier}{suffix}").status_code == 404
    assert client.get("/api/v1/report-jobs").json()["items"] == []
    client.cookies.clear()
    assert client.get(result["download_url"]).status_code == 401


def test_tampered_source_fails_without_rendering(reports, snapshot, monkeypatch):
    client, factory = reports
    with factory() as session:
        job = report_jobs.enqueue(session, "report-owner", snapshot, "xlsx")
        job.snapshot_hash = "0" * 64
        session.commit()
        identifier = job.id
    monkeypatch.setattr(report_jobs, "render", lambda *_: pytest.fail("Tampered source rendered"))
    assert report_jobs.execute(identifier)
    result = client.get(f"/api/v1/report-jobs/{identifier}").json()
    assert result["status"] == "FAILED"
    assert result["download_url"] is None
    assert client.get(result["source_url"]).status_code == 409


def test_tampered_output_and_expired_retention_are_rejected(reports, snapshot):
    client, factory = reports
    with factory() as session:
        job = report_jobs.enqueue(session, "report-owner", snapshot, "xlsx")
        session.commit()
        identifier = job.id
    report_jobs.execute(identifier)
    with factory() as session:
        job = session.get(ReportJob, identifier)
        ObjectStorage().put_bytes(key=job.object_key, data=b"tampered", content_type="text/plain")
    assert client.get(f"/api/v1/report-jobs/{identifier}/download").status_code == 409
    with factory() as session:
        session.get(ReportJob, identifier).expires_at = datetime.now(UTC) - timedelta(seconds=1)
        session.commit()
    assert client.get(f"/api/v1/report-jobs/{identifier}/download").status_code == 410


def test_active_job_limit(reports, snapshot, monkeypatch):
    client, _ = reports
    monkeypatch.setattr(report_api, "capture", lambda *_: snapshot)
    for _ in range(4):
        assert client.post("/api/v1/report-jobs", json={"kind": "portfolio"}).status_code == 202
    assert client.post("/api/v1/report-jobs", json={"kind": "portfolio"}).status_code == 429


def test_real_portfolio_source_is_explicitly_dated(ledger_session: Session) -> None:
    from datetime import date

    snapshot = capture(
        ledger_session, ReportRequest(kind="portfolio", portfolio="book", as_of=date(2026, 1, 9))
    )
    assert snapshot.references["valuation_run_id"]
    assert snapshot.references["portfolio_id"] == "book"
    assert snapshot.payload["portfolio"]["id"] == "book"
    assert snapshot.calculation_version != "UNVERSIONED"


def test_missing_analysis_does_not_fabricate_report(ledger_session: Session) -> None:
    with pytest.raises(ValueError, match="completed matching"):
        capture(ledger_session, ReportRequest(kind="backtest"))


def test_expired_worker_cannot_publish_after_lease_is_lost(reports, snapshot, monkeypatch):
    _, factory = reports
    with factory() as session:
        job = report_jobs.enqueue(session, "report-owner", snapshot, "xlsx")
        session.commit()
        identifier = job.id

    def delayed_render(*args):
        with factory() as session:
            session.get(ReportJob, identifier).status = "FAILED"
            session.commit()
        return render(*args)

    monkeypatch.setattr(report_jobs, "render", delayed_render)
    assert not report_jobs.execute(identifier)
    with factory() as session:
        assert session.get(ReportJob, identifier).status == "FAILED"
        assert session.get(ReportJob, identifier).content_hash is None


def test_saved_comparables_and_macro_capture(ledger_session: Session) -> None:
    run = models.AnalysisRun(
        kind="comparables",
        name="Saved peers",
        status="SUCCEEDED",
        result={"source": "TEST", "quality": "FILE", "peers": [{"symbol": "AAA"}]},
        parameters={},
    )
    ledger_session.add(run)
    ledger_session.flush()
    result = capture(ledger_session, ReportRequest(kind="comps", analysis_run_id=run.id))
    assert result.references["analysis_run_id"] == run.id
    macro = capture(ledger_session, ReportRequest(kind="macro"))
    assert macro.payload["items"] == []
    assert macro.data_as_of is None


def test_expired_worker_lease_is_failed_not_automatically_retried(reports, snapshot):
    _, factory = reports
    with factory() as session:
        job = report_jobs.enqueue(session, "report-owner", snapshot, "xlsx")
        job.status, job.started_at = "RUNNING", datetime.now(UTC) - timedelta(minutes=16)
        session.commit()
        identifier = job.id
    report_jobs.run_once()
    with factory() as session:
        job = session.get(ReportJob, identifier)
        assert job.status == "FAILED"
        assert job.error_category == "WorkerLeaseExpired"
