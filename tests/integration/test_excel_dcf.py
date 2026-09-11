"""Opt-in Excel recalculation of generated fictional workbooks, never user files."""

import hashlib
import os
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "services/api"))

from app.dcf_workbook import Formula, model_sheets  # noqa: E402
from app.equity_valuation import DcfRequest, dcf_scenarios  # noqa: E402
from app.report_contracts import ReportSnapshot  # noqa: E402
from app.report_render import render  # noqa: E402

pytestmark = pytest.mark.skipif(
    os.environ.get("KNK_EXCEL_INTEGRATION") != "1" or os.name != "nt",
    reason="Requires explicit opt-in and installed Windows Excel; not a mocked calculation.",
)


@pytest.mark.parametrize(
    "method,years,price,calculated_wacc",
    [
        ("PERPETUITY", 1, "12", False),
        ("PERPETUITY", 5, None, True),
        ("EXIT_MULTIPLE", 10, "8", False),
    ],
)
def test_excel_recalculation_matches_every_cached_formula(
    method: str,
    years: int,
    price: str | None,
    calculated_wacc: bool,
) -> None:
    request = DcfRequest.model_validate(
        {
            "symbol": "FICTIONAL",
            "initial_working_capital": 0,
            "debt_override": "25",
            "cash_override": "12",
            "shares_override": "11",
            "wacc_inputs": {} if calculated_wacc else None,
            "apply_calculated_wacc": calculated_wacc,
            "scenarios": [
                {
                    "name": "BASE",
                    "revenue_growth": [".08"] * years,
                    "ebit_margins": [".25"] * years,
                    "terminal_method": method,
                },
                {
                    "name": "BEAR",
                    "revenue_growth": ["-.08"] * years,
                    "ebit_margins": ["-.1"] * years,
                    "terminal_method": method,
                },
            ],
        }
    )
    payload = dcf_scenarios(
        {"revenue": "100.1234567", "debt": "20", "cash": "10", "shares": "10"}, request, price
    )
    payload.update(
        parameters=request.model_dump(mode="json"), quote={"price": price}, unit="millions"
    )
    snapshot = ReportSnapshot(
        title="KnK Capital | Fictional DCF validation",
        kind="dcf",
        requested_at="2026-01-10T00:00:00Z",
        data_as_of="2026-01-09T00:00:00Z",
        source="Fictional test statements",
        quality="DEMO",
        currency="USD",
        calculation_version="knk-fcff-1.0",
        references={"analysis_run_id": "isolated-excel-fixture"},
        sections=[],
        payload=payload,
    )
    directory = ROOT / "logs" / "dcf-validation" / uuid4().hex
    directory.mkdir(parents=True)
    original = render(snapshot, "xlsx")
    source = directory / "source.xlsx"
    source.write_bytes(original)
    recalculate(source, directory / "baseline")
    actual = load_workbook(directory / "baseline/recalculated.xlsx", data_only=True)
    checked = 0
    for sheet in model_sheets(snapshot):
        for cell in sheet.cells:
            if isinstance(cell.value, Formula):
                value = actual[sheet.name][cell.address].value
                if isinstance(cell.value.cached, str):
                    assert value == cell.value.cached, (sheet.name, cell.address, value)
                else:
                    assert value == pytest.approx(cell.value.cached, rel=1e-10, abs=1e-10), (
                        sheet.name,
                        cell.address,
                        value,
                    )
                checked += 1
    assert checked >= 200
    assert source.read_bytes() == original
    changed = load_workbook(BytesIO(original))
    changed["DCF Inputs"]["B7"] = 125
    changed.save(directory / "changed.xlsx")
    recalculate(directory / "changed.xlsx", directory / "changed")
    recomputed = load_workbook(directory / "changed/recalculated.xlsx", data_only=True)
    assert recomputed["DCF BASE"]["C10"].value == pytest.approx(135)
    assert recomputed["DCF Summary"]["H2"].value != 0
    invalid = load_workbook(BytesIO(original))
    invalid["DCF BASE"]["H6"] = "PERPETUITY"
    invalid["DCF BASE"]["B6"] = 0.02
    invalid["DCF BASE"]["D6"] = 0.03
    invalid.save(directory / "invalid.xlsx")
    recalculate(directory / "invalid.xlsx", directory / "invalid")
    rejected = load_workbook(directory / "invalid/recalculated.xlsx", data_only=True)
    assert rejected["DCF BASE"]["B31"].value == "#N/A"
    print(
        f"Excel verified {checked} formulas; baseline SHA256={hashlib.sha256(original).hexdigest()}; evidence={directory}"
    )


def recalculate(source: Path, output: Path) -> None:
    result = subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(ROOT / "scripts/recalculate-excel.ps1"),
            "-InputPath",
            str(source),
            "-OutputDirectory",
            str(output),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
        creationflags=subprocess.CREATE_NO_WINDOW,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert (output / "rendered.pdf").read_bytes().startswith(b"%PDF")
